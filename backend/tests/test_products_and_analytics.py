import io
from decimal import Decimal
from urllib.parse import urlsplit

from openpyxl import load_workbook
from sqlalchemy import select, text, update

from app.core.db import session_scope
from app.models import Listing, Notification
from tests.conftest import make_user, run_keyword_job


def test_search_filters_and_cursor_pagination(client, admin):
    h = admin["headers"]
    run_keyword_job(client, h, keyword="yoga", limit=60)
    seen, cursor = [], None
    while True:
        url = "/api/v1/products?limit=17&sort=price_asc" + (f"&cursor={cursor}" if cursor else "")
        page = client.get(url, headers=h).json()
        seen += page["items"]
        cursor = page["next_cursor"]
        if not cursor:
            break
    ids = [i["id"] for i in seen]
    assert len(ids) == len(set(ids))  # stable keyset pagination: no duplicates, nothing skipped
    total = client.get("/api/v1/products?with_total=true&limit=1", headers=h).json()["total"]
    assert len(ids) == total
    prices = [Decimal(i["current_price"]) for i in seen if i["current_price"] is not None]
    assert prices == sorted(prices)

    filtered = client.get("/api/v1/products?price_min=10&price_max=40&stock_status=in_stock", headers=h).json()["items"]
    assert all(10 <= Decimal(i["current_price"]) <= 40 and i["stock_status"] == "in_stock" for i in filtered)
    fts = client.get("/api/v1/products?q=yoga", headers=h).json()["items"]
    assert fts and all("yoga" in i["title"].lower() for i in fts)
    updated_cursor = client.get("/api/v1/products?limit=1&sort=updated", headers=h).json()["next_cursor"]
    bad = client.get(f"/api/v1/products?sort=title&cursor={updated_cursor}", headers=h)
    assert bad.status_code == 400


def test_product_detail_history_metrics_sources(client, admin):
    h = admin["headers"]
    run_keyword_job(client, h, keyword="mat", limit=5)
    pid = client.get("/api/v1/products?limit=1", headers=h).json()["items"][0]["id"]
    d = client.get(f"/api/v1/products/{pid}", headers=h).json()
    assert d["title"].startswith("[DEMO]") and d["gtin"] and d["images"]
    hist = client.get(f"/api/v1/products/{pid}/history", headers=h).json()
    assert len(hist) == 1 and hist[0]["price"] == d["current_price"]
    metrics = {m["metric_name"]: m for m in client.get(f"/api/v1/products/{pid}/metrics", headers=h).json()}
    assert metrics["competitor_sales"]["measurement_type"] == "unavailable"
    assert metrics["competitor_sales"]["numeric_value"] is None
    assert metrics["estimated_monthly_units"]["measurement_type"] == "estimated"
    sources = {s["field_name"]: s for s in client.get(f"/api/v1/products/{pid}/sources", headers=h).json()}
    assert sources["current_price"]["connector"] == "demo"


def test_profit_and_opportunity_for_listing(client, admin):
    h = admin["headers"]
    run_keyword_job(client, h, keyword="mat", limit=5)
    item = client.get("/api/v1/products?limit=1", headers=h).json()["items"][0]
    pid = item["id"]
    before = client.get(f"/api/v1/products/{pid}/profit", headers=h).json()
    assert not before["is_complete"] and "supplier_cost" in before["missing_inputs"]
    opp_before = client.get(f"/api/v1/products/{pid}/opportunity", headers=h).json()

    r = client.post(
        "/api/v1/fee-profiles",
        headers=h,
        json={
            "marketplace_code": "DEMO_UK",
            "name": "Demo standard",
            "referral_rate": "0.15",
            "fixed_fee": "0.30",
            "currency_code": "GBP",
        },
    )
    assert r.status_code == 201, r.text
    r = client.post(
        f"/api/v1/products/{pid}/costs",
        headers=h,
        json={
            "currency_code": "USD",
            "supplier_cost": "4.00",
            "shipping_cost": "1.00",
            "assessment": {"supplier_reliability": 80, "delivery_speed": 60, "compliance_risk": 10},
        },
    )
    assert r.status_code == 201, r.text
    # cost is in USD; without an exchange rate the result must say so rather than guess
    p = client.get(f"/api/v1/products/{pid}/profit", headers=h).json()
    assert "exchange_rate" in p["missing_inputs"]
    client.post(
        "/api/v1/exchange-rates",
        headers=h,
        json={"base_currency": "GBP", "quote_currency": "USD", "rate": "1.25", "effective_at": "2026-01-01T00:00:00Z"},
    )
    p = client.get(f"/api/v1/products/{pid}/profit?target_margin_pct=30", headers=h).json()
    assert p["is_complete"], p
    assert p["fx"][0]["from"] == "USD"
    assert Decimal(p["net_revenue"]) == (Decimal(item["current_price"]) / Decimal("1.2")).quantize(Decimal("0.01"))
    assert p["target_margin_price"] is not None
    opp = client.get(f"/api/v1/products/{pid}/opportunity", headers=h).json()
    assert opp["confidence"] > opp_before["confidence"]
    assert "profit_margin" not in opp["missing"]

    calc = client.post(
        "/api/v1/analytics/profit/calculate",
        headers=h,
        json={"selling_price": "20", "supplier_cost": "5", "fees": {"referral_rate": "0.1"}},
    ).json()
    assert calc["profit"] == "13.00"


def test_overview_and_opportunities(client, admin):
    h = admin["headers"]
    run_keyword_job(client, h, keyword="yoga", limit=20)
    run_keyword_job(client, h, keyword="yoga", limit=20, marketplace="DEMO_US")
    o = client.get("/api/v1/analytics/overview", headers=h).json()
    assert {m["marketplace"] for m in o["by_marketplace"]} == {"DEMO_UK", "DEMO_US"}
    assert o["total_listings"] > 0 and o["listings_with_metrics"]["estimated"] > 0
    opps = client.get("/api/v1/analytics/opportunities?limit=5", headers=h).json()
    scores = [i["score"] for i in opps["items"]]
    assert scores == sorted(scores, reverse=True) and all(0 <= s <= 100 for s in scores)


def test_gtin_matching_and_review(client, admin):
    h = admin["headers"]
    run_keyword_job(client, h, keyword="yoga", limit=20, marketplace="DEMO_UK")
    run_keyword_job(client, h, keyword="yoga", limit=20, marketplace="DEMO_US")
    matches = client.get("/api/v1/matches", headers=h).json()["items"]
    assert matches and all(m["match_method"] == "gtin_exact" for m in matches)
    m = matches[0]
    assert m["source"]["marketplace_code"] != m["target"]["marketplace_code"]
    r = client.post(f"/api/v1/matches/{m['id']}/review", headers=h, json={"decision": "confirmed"})
    assert r.status_code == 200 and r.json()["status"] == "confirmed"
    a = client.get(f"/api/v1/products/{m['source_listing_id']}", headers=h).json()
    b = client.get(f"/api/v1/products/{m['target_listing_id']}", headers=h).json()
    assert a["canonical_product_id"] and a["canonical_product_id"] == b["canonical_product_id"]
    assert a["id"] != b["id"]  # listings stay distinct
    cmp = client.post(
        "/api/v1/products/compare", headers=h, json={"listing_ids": [a["id"], b["id"]], "currency": "GBP"}
    ).json()
    assert len(cmp["items"]) == 2
    # fuzzy suggestions run without error and never overwrite reviewed decisions
    client.post(f"/api/v1/products/{a['id']}/matches/suggest", headers=h)
    assert client.get("/api/v1/matches?status=confirmed", headers=h).json()["items"][0]["id"] == m["id"]


def test_price_drop_alert(client, admin):
    h = admin["headers"]
    run_keyword_job(client, h, keyword="mat", limit=5)
    pid = client.get("/api/v1/products?limit=1", headers=h).json()["items"][0]["id"]
    r = client.post(
        "/api/v1/alert-rules", headers=h, json={"rule_type": "price_drop_pct", "threshold": "5", "listing_id": pid}
    )
    assert r.status_code == 201
    with session_scope() as db:  # make the stored price higher so the next refresh is a drop
        db.execute(
            update(Listing)
            .where(Listing.id == pid)
            .values(current_price=Listing.current_price * 2, content_hash="stale")
        )
    run_keyword_job(client, h, keyword="mat", limit=5)
    notes = client.get("/api/v1/notifications", headers=h).json()
    assert any("Price dropped" in n["title"] for n in notes)
    with session_scope() as db:
        assert db.scalar(select(Notification.data).limit(1))["listing_id"] == pid


def test_exports_xlsx_and_csv(client, admin):
    h = admin["headers"]
    run_keyword_job(client, h, keyword="yoga", limit=15)
    with session_scope() as db:  # inject a malicious title to prove formula neutralisation
        db.execute(
            text(
                'UPDATE marketplace_listings SET title = \'=HYPERLINK("http://evil","x")\' '
                "WHERE id = (SELECT min(id) FROM marketplace_listings)"
            )
        )
    r = client.post("/api/v1/exports", headers=h, json={"export_format": "xlsx"})
    assert r.status_code == 202, r.text
    exp = client.get(f"/api/v1/exports/{r.json()['id']}", headers=h).json()
    assert exp["export_status"] == "completed", exp
    assert exp["manifest"]["sheets"]["Products"] == exp["row_count"] > 0
    link = client.post(f"/api/v1/exports/{exp['id']}/link", headers=h).json()
    parts = urlsplit(link["url"])
    dl = client.get(f"{parts.path}?{parts.query}")
    assert dl.status_code == 200
    wb = load_workbook(io.BytesIO(dl.content), read_only=True)
    assert wb.sheetnames == [
        "Products",
        "Variations",
        "Images",
        "Price History",
        "Metrics",
        "Profit Analysis",
        "Manifest",
    ]
    titles = [row[8] for row in wb["Products"].iter_rows(min_row=2, values_only=True)]
    assert '\'=HYPERLINK("http://evil","x")' in titles
    metric_types = {row[4] for row in wb["Metrics"].iter_rows(min_row=2, values_only=True)}
    assert {"actual", "estimated", "unavailable"} <= metric_types
    # tampered or foreign tokens are rejected
    tampered = parts.query[:-1] + ("0" if parts.query[-1] != "0" else "1")
    assert client.get(f"{parts.path}?{tampered}").status_code == 403
    other_user = parts.query.replace("u=", "u=0", 1)
    assert client.get(f"{parts.path}?{other_user}").status_code in (403, 422)

    r = client.post("/api/v1/exports", headers=h, json={"export_format": "csv", "filters": {"q": "yoga"}})
    csv_exp = client.get(f"/api/v1/exports/{r.json()['id']}", headers=h).json()
    assert csv_exp["export_status"] == "completed" and csv_exp["row_count"] > 0


def test_viewer_cannot_export_or_see_others_exports(client, admin):
    viewer = make_user(client, admin["headers"], "v@example.com", "viewer")
    assert client.post("/api/v1/exports", headers=viewer, json={}).status_code == 403
    r = client.post("/api/v1/exports", headers=admin["headers"], json={"export_format": "csv"})
    assert client.get(f"/api/v1/exports/{r.json()['id']}", headers=viewer).status_code == 403


def test_pagination_includes_rows_without_sort_value(client, admin):
    h = admin["headers"]
    run_keyword_job(client, h, keyword="yoga", limit=30)
    with session_scope() as db:
        db.execute(text("UPDATE marketplace_listings SET current_price = NULL, sales_rank = NULL WHERE id % 3 = 0"))
    for sort in ("price_asc", "price_desc", "rank", "title", "updated"):
        seen, cursor = [], None
        while True:
            page = client.get(
                f"/api/v1/products?limit=7&sort={sort}" + (f"&cursor={cursor}" if cursor else ""), headers=h
            ).json()
            seen += [i["id"] for i in page["items"]]
            cursor = page["next_cursor"]
            if not cursor:
                break
        total = client.get("/api/v1/products?with_total=true&limit=1", headers=h).json()["total"]
        assert len(seen) == len(set(seen)) == total, sort


def test_overview_uses_materialised_view_for_large_orgs(client, admin, monkeypatch):
    from app.services import analytics

    h = admin["headers"]
    run_keyword_job(client, h, keyword="yoga", limit=20)
    live = client.get("/api/v1/analytics/overview", headers=h).json()
    assert live["aggregates_as_of"] is None
    with session_scope() as db:
        analytics.refresh_summary_view(db)
    monkeypatch.setattr(analytics, "MV_THRESHOLD", 1)
    pre = client.get("/api/v1/analytics/overview", headers=h).json()
    assert pre["aggregates_as_of"] is not None
    assert pre["total_listings"] == live["total_listings"]
    assert pre["stock_status"] == live["stock_status"]
